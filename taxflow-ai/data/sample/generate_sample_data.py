#!/usr/bin/env python3
"""Generate sample test data for TaxFlow AI.

Creates sample PDF invoices and Excel files for testing.
Run: python data/sample/generate_sample_data.py
"""

import os
import csv
from pathlib import Path

SAMPLE_DIR = Path(__file__).resolve().parent


def create_invoice_pdf(text_file: str, output_pdf: str):
    """Create a PDF from the text template."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet

        input_path = SAMPLE_DIR / text_file
        output_path = SAMPLE_DIR / output_pdf

        if not input_path.exists():
            print(f"  ⚠️  Template {text_file} not found, skipping")
            return

        with open(input_path, "r", encoding="utf-8") as f:
            content = f.read()

        doc = SimpleDocTemplate(str(output_path), pagesize=A4,
                                rightMargin=50, leftMargin=50,
                                topMargin=50, bottomMargin=50)
        styles = getSampleStyleSheet()

        elements = []
        for line in content.split("\n"):
            if line.strip():
                elements.append(Paragraph(line.replace("\n", "<br/>"), styles["Normal"]))
            elements.append(Spacer(1, 6))

        doc.build(elements)
        print(f"  ✅ Created {output_pdf}")

    except ImportError:
        print(f"  ⚠️  ReportLab not installed. Creating text version only.")
        # Text files are already present as source
        print(f"  📝 Using text template: {text_file}")


def create_sample_excel():
    """Create a sample Excel invoice."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

        output_path = SAMPLE_DIR / "invoice_3.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Export Invoice"

        # Title
        ws.merge_cells("A1:F1")
        ws.cell(row=1, column=1, value="EXPORT INVOICE")
        ws.cell(row=1, column=1).font = Font(bold=True, size=16, color="1F4E79")

        # Headers
        headers = ["#", "HSN Code", "Description", "Quantity", "Rate ($)", "Amount ($)"]
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Data
        data = [
            [1, "5208", "Cotton Fabric 100%", "5000m", 5.00, 25000.00],
        ]

        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        for row_idx, row_data in enumerate(data, 4):
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx in (5, 6):
                    cell.number_format = '$ #,##0.00'
                cell.alignment = Alignment(horizontal="right" if col_idx > 1 else "center")

        # Total
        ws.cell(row=6, column=5, value="Total FOB:")
        ws.cell(row=6, column=5).font = Font(bold=True)
        ws.cell(row=6, column=6, value=25000.00)
        ws.cell(row=6, column=6).number_format = '$ #,##0.00'
        ws.cell(row=6, column=6).font = Font(bold=True)

        # Invoice details
        details = [
            ("Invoice No:", "EXP-2024-001"),
            ("Date:", "25-Jan-2024"),
            ("Exporter:", "ExportHouse Inc, Chennai"),
            ("GSTIN:", "33AAAE7890L1Z9"),
            ("Buyer:", "International Buyer LLC, New York, USA"),
            ("IEC:", "1234567890"),
        ]

        for i, (label, value) in enumerate(details):
            ws.cell(row=9 + i, column=1, value=label).font = Font(bold=True)
            ws.cell(row=9 + i, column=2, value=value)

        for col in range(1, 7):
            ws.column_dimensions[chr(64 + col)].width = 18

        wb.save(str(output_path))
        print(f"  ✅ Created {output_pdf}")

    except ImportError:
        # Fallback: create CSV
        output_path = SAMPLE_DIR / "invoice_3.csv"
        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["#", "HSN Code", "Description", "Quantity", "Rate ($)", "Amount ($)"])
            writer.writerow([1, "5208", "Cotton Fabric 100%", "5000m", 5.00, 25000.00])
        print(f"  ✅ Created invoice_3.csv (Excel not available)")


def verify_sample_data():
    """Check all sample files exist."""
    expected_files = [
        "invoice_1.txt", "invoice_2.txt", "invoice_3.txt",
        "invoice_1.pdf", "invoice_2.pdf", "invoice_3.xlsx",
        "bank_statement.csv",
    ]

    print("\n📋 Sample Data Files:")
    for f in expected_files:
        path = SAMPLE_DIR / f
        if path.exists():
            size = path.stat().st_size
            print(f"  ✅ {f:30s} ({size:,} bytes)")
        else:
            print(f"  ⚠️  {f:30s} NOT FOUND")


def main():
    """Generate all sample data."""
    print("=" * 60)
    print("  📦 TaxFlow AI - Sample Data Generator")
    print("=" * 60)

    print("\n📄 Generating PDF invoices from text templates...")
    create_invoice_pdf("invoice_1.txt", "invoice_1.pdf")
    create_invoice_pdf("invoice_2.txt", "invoice_2.pdf")

    print("\n📊 Generating Excel invoice...")
    create_sample_excel()

    print("\n✅ Sample data generation complete!")

    # Show what was created
    verify_sample_data()

    print("\n" + "=" * 60)
    print("  🚀 Sample data is ready! Upload these files to TaxFlow AI.")
    print("  💡 Default login: admin / admin123")
    print("=" * 60)


if __name__ == "__main__":
    main()
