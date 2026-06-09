"""Document processing utilities shared across extractors."""

import os
import tempfile
from pathlib import Path
from typing import Any

from app.config import settings


def get_secure_temp_path(filename: str) -> str:
    """Get a secure temporary path for file processing."""
    temp_dir = Path(tempfile.mkdtemp(prefix="taxflow_"))
    return str(temp_dir / filename)


def cleanup_temp(paths: list[str]):
    """Clean up temporary files."""
    for p in paths:
        try:
            if os.path.isfile(p):
                os.remove(p)
            elif os.path.isdir(p):
                import shutil
                shutil.rmtree(p, ignore_errors=True)
        except Exception:
            pass


def validate_file_extension(filename: str) -> bool:
    """Check if file extension is supported."""
    ext = Path(filename).suffix.lower()
    return ext in settings.SUPPORTED_EXTENSIONS


def format_currency(amount: float) -> str:
    """Format amount as Indian currency."""
    if amount >= 10000000:
        return f"₹{amount / 10000000:.2f} Cr"
    elif amount >= 100000:
        return f"₹{amount / 100000:.2f} L"
    return f"₹{amount:,.2f}"


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF file using multiple strategies."""
    import pdfplumber
    import fitz  # PyMuPDF

    text_parts = []

    # Strategy 1: pdfplumber (good for structured tables)
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        text_parts.append(" | ".join(str(cell or "") for cell in row))
    except Exception:
        pass

    # Strategy 2: PyMuPDF (good for general text extraction)
    try:
        doc = fitz.open(file_path)
        for page in doc:
            text_parts.append(page.get_text())
        doc.close()
    except Exception:
        pass

    return "\n".join(text_parts)


def extract_text_from_xlsx(file_path: str) -> dict[str, Any]:
    """Extract text and structure from an Excel file."""
    import openpyxl

    wb = openpyxl.load_workbook(file_path, data_only=True)
    result = {"sheets": {}}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(cell) if cell is not None else "" for cell in row])
        result["sheets"][sheet_name] = {
            "headers": rows[0] if rows else [],
            "data": rows[1:] if len(rows) > 1 else [],
            "total_rows": len(rows),
        }
    wb.close()
    return result


def extract_text_from_csv(file_path: str) -> list[dict[str, str]]:
    """Extract text from a CSV file."""
    import csv

    rows = []
    with open(file_path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append({k.strip(): v.strip() for k, v in row.items()})
    return rows
